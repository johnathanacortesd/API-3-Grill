# ======================================
# Subtema collages, PKL batching, progress, row counts
# ======================================
import io
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from ai_analyzer import (
    cluster_similar_rows,
    enrich_rows_with_ai,
    ensure_subtema_distinct_from_tema,
)
from app import PIPELINE_STEPS, _active_step
from pipeline import BASE_OUTPUT_COLUMNS, KEY_MAP, generate_output_excel, process_dossier
from test_pkl_classifier import _dump_bytes, _theme_pipeline, _mini_xlsx_bytes
from test_pkl_subtema_grouping import _PredictByKeyword, _news_row


class _SpyPredict:
    def __init__(self, inner):
        self.inner = inner
        self.calls = []

    def predict(self, texts):
        batch = list(texts)
        self.calls.append(len(batch))
        return self.inner.predict(batch)


class SubtemaCollageTests(unittest.TestCase):
    def test_rejects_de_de_and_keyword_collage(self):
        ctx = (
            "De ese barrio salió, primero, un joven que se hizo abogado "
            "en la Universidad Simón Bolívar de Barranquilla."
        )
        brand = "Universidad Simón Bolívar"
        title = "Historia de un egresado"
        for raw in (
            "Abogado de universidad de barranquilla de joven",
            "Formación de de universidad de",
            "Universidad de de Antioquia",
        ):
            sub = ensure_subtema_distinct_from_tema("Mención", raw, brand, title, ctx)
            low = sub.strip().lower()
            self.assertNotIn(" de de ", f" {low} ")
            self.assertNotEqual(low, raw.lower())
            self.assertGreaterEqual(len(sub.split()), 4)
            self.assertLessEqual(len(sub.split()), 7)
            self.assertFalse(low.startswith("abogado de universidad de"))
            self.assertIn("simón bolívar", low)

    def test_keeps_grammatical_two_de_noun_phrase(self):
        sub = ensure_subtema_distinct_from_tema(
            "Eventos o Proyectos",
            "Apertura de sede norte de Cali",
            "UdeA",
            "Inauguran sede norte",
            "La universidad inauguró una sede norte de Cali para ampliar cobertura.",
        )
        self.assertEqual(sub.lower(), "apertura de sede norte de cali")

    def test_keeps_specific_llm_subtema_on_pkl_tema_path(self):
        sub = ensure_subtema_distinct_from_tema(
            "Mención",
            "Apertura de sede norte",
            "UdeA",
            "Inauguran sede norte",
            "La universidad inauguró una sede para ampliar cobertura educativa en el norte.",
        )
        self.assertEqual(sub, "Apertura de sede norte")


class PklBatchAndProgressTests(unittest.TestCase):
    def test_theme_pkl_predicts_once_per_axis_not_per_cluster(self):
        rows = [
            _news_row("Inauguración sede norte en Cali", "la universidad inauguró sede norte"),
            _news_row("Denuncian cobros excesivos en peajes", "usuarios reportan tarifas altas en peajes"),
            _news_row("Feria de becas de posgrado", "abren convocatoria de becas de posgrado"),
            _news_row("Inauguración sede norte en Cali hoy", "inauguró la sede norte en cali"),
        ]
        spy = _SpyPredict(_PredictByKeyword([], "Mención"))
        logs = []
        with patch("ai_analyzer.OpenAI"):
            with patch("ai_analyzer._call_openai_cluster") as mock_llm:
                mock_llm.return_value = ("Neutro", "Educación Superior", "Apertura de sede norte")
                out = enrich_rows_with_ai(
                    rows,
                    KEY_MAP,
                    "UdeA",
                    [],
                    "sk-test",
                    theme_model=spy,
                    progress_callback=lambda pct, msg: logs.append((pct, msg)),
                )
        self.assertEqual(spy.calls, [3], "similar inauguraciones share a cluster; 3 hechos → 1 lote")
        self.assertEqual(len(out), 4)
        self.assertEqual(out[0]["Tema_IA"], out[3]["Tema_IA"])
        self.assertEqual(out[0]["Subtema_IA"], out[3]["Subtema_IA"])
        joined = " | ".join(m for _, m in logs)
        self.assertIn("Temas listos", joined)
        self.assertIn("lote", joined.lower())
        self.assertIn("hechos", joined.lower())
        self.assertTrue(any("Agrupando" in m for _, m in logs))
        self.assertTrue(any("Etiquetando" in m for _, m in logs))

    def test_ai_path_does_not_inflate_rows(self):
        rows = [
            _news_row("Nota uno", "contexto uno de la universidad"),
            _news_row("Nota dos", "contexto dos de la universidad"),
            {
                **_news_row("Nota dup", "n/a"),
                "is_duplicate": True,
                "Subtema_IA": "-",
            },
        ]
        with patch("ai_analyzer.OpenAI"):
            with patch("ai_analyzer._call_openai_cluster") as mock_llm:
                mock_llm.return_value = ("Neutro", "Gestión Institucional", "Hecho informativo local")
                out = enrich_rows_with_ai(rows, KEY_MAP, "UdeA", [], "sk-test")
        self.assertEqual(len(out), 3)
        self.assertEqual(out[2]["Tono_IA"], "Duplicada")


class ClusteringAndExportTests(unittest.TestCase):
    def test_similar_titles_cluster_dissimilar_do_not(self):
        rows = [
            {
                "Título": "Inauguración sede norte en Cali",
                "Resumen - Aclaracion": "la universidad inauguró sede norte",
                "is_duplicate": False,
            },
            {
                "Título": "Inauguración sede norte en Cali hoy",
                "Resumen - Aclaracion": "inauguró la sede norte",
                "is_duplicate": False,
            },
            {
                "Título": "Denuncian cobros excesivos en peajes",
                "Resumen - Aclaracion": "usuarios quejan tarifas de peaje en la vía",
                "is_duplicate": False,
            },
        ]
        cm = cluster_similar_rows(rows, KEY_MAP, [])
        self.assertEqual(cm[0], cm[1])
        self.assertNotEqual(cm[0], cm[2])

    def test_export_with_ai_columns_keeps_row_count_and_black_links(self):
        rows = [
            {
                "ID Noticia": 101,
                "Título": "Nota A",
                "revalorización": 10,
                "resumen corto": "texto",
                "Link Nota": {"value": "Link", "url": "https://example.com/a"},
                "Link (Streaming - Imagen)": {"value": "Link", "url": "https://example.com/sa"},
                "Contexto analizado": "ctx",
                "Tono_IA": "Neutro",
                "Tema_IA": "Mención",
                "Subtema_IA": "Apertura de sede norte",
            },
            {
                "ID Noticia": 102,
                "Título": "Nota B",
                "revalorización": 20,
                "resumen corto": "texto",
                "Link Nota": {"value": "Link", "url": "https://example.com/b"},
                "Contexto analizado": "ctx",
                "Tono_IA": "Positivo",
                "Tema_IA": "Entrevista",
                "Subtema_IA": "Diálogo con el rector",
            },
        ]
        rev_idx = BASE_OUTPUT_COLUMNS.index("revalorización")
        ai_cols = ["Contexto analizado", "Tono_IA", "Tema_IA", "Subtema_IA"]
        cols = BASE_OUTPUT_COLUMNS[: rev_idx + 1] + ai_cols + BASE_OUTPUT_COLUMNS[rev_idx + 1 :]
        data = generate_output_excel(rows, KEY_MAP, columns_to_use=cols)
        wb = load_workbook(io.BytesIO(data))
        ws = wb["Resultado"]
        self.assertEqual(ws.max_row, 3)
        headers = [c.value for c in ws[1]]
        self.assertEqual(headers, cols)
        self.assertIn("Link Nota", headers)
        self.assertIn("Subtema_IA", headers)
        link_cell = ws.cell(row=2, column=headers.index("Link Nota") + 1)
        self.assertIn("HYPERLINK(", str(link_cell.value or "").upper())
        self.assertTrue(str(getattr(link_cell.font.color, "rgb", "FF000000")).upper().endswith("000000"))

    def test_pkl_pipeline_does_not_add_rows(self):
        result = process_dossier(
            _mini_xlsx_bytes(),
            {"eltiempo": "Nacional"},
            {"eltiempo": "El Tiempo"},
            ai_config={
                "enabled": False,
                "brand": "Marca Demo",
                "aliases": [],
                "tone_pkl_bytes": None,
                "theme_pkl_bytes": _dump_bytes(_theme_pipeline()),
            },
        )
        self.assertEqual(result["total_rows"], 2)


class ProgressStepTests(unittest.TestCase):
    def test_pipeline_steps_include_group_and_label(self):
        keys = [k for k, _ in PIPELINE_STEPS]
        self.assertEqual(keys, ["config", "read", "norm", "dups", "group", "label", "export"])

    def test_active_step_maps_grouping_and_pkl_labeling(self):
        self.assertEqual(_active_step(74, "Agrupando noticias similares… 40/200 notas, 12 hechos"), "group")
        self.assertEqual(_active_step(76, "Temas listos: 120 hechos únicos de 400 notas. Clasificando PKL en lote…"), "label")
        self.assertEqual(_active_step(88, "Etiquetando con IA… 15/120 hechos"), "label")
        self.assertEqual(_active_step(94, "Generando archivo de resultado… 1/10 filas"), "export")


if __name__ == "__main__":
    unittest.main()

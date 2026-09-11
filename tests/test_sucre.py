# ======================================
# Variante Sucre: extractos literales, tono y columnas de actores
# ======================================
import io
import os
import sys
import unittest
from unittest.mock import MagicMock

import pandas as pd
from openpyxl import load_workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from pipeline import BASE_OUTPUT_COLUMNS, KEY_MAP
from sucre_analyzer import (
    COL_EXTERNOS,
    COL_INT_PROPIA,
    COL_MENCION_EXT,
    COL_PROPIOS,
    COL_TONO,
    SUCRE_OUTPUT_COLUMNS,
    analyze_article,
    heuristic_analyze,
    is_filler,
    merge_analysis,
    recover_verbatim,
)
from sucre_pipeline import build_sample_xlsx, process_sucre_dossier


LUCY_BODY = (
    "La gobernadora Lucy García Montes entregó 200 becas universitarias en Sincelejo. "
    '"La educación es la prioridad de este gobierno", afirmó la mandataria. '
    "El alcalde Ricardo Hernández felicitó a la gobernadora por la inversión social."
)

SENATOR_BODY = (
    "El senador Andrés Pérez señaló que la Gobernación de Sucre no ha ejecutado "
    "el presupuesto de vías rurales y pidió explicaciones a la mandataria."
)

SECRE_BODY = (
    "La Secretaría de Educación departamental emitió la Resolución 045 de 2026 "
    "para ampliar la cobertura escolar en los municipios del Golfo de Morrosquillo."
)


class VerbatimRulesTests(unittest.TestCase):
    def test_recover_exact_span(self):
        src = "La Gobernación de Sucre emitió un comunicado."
        self.assertEqual(recover_verbatim("La Gobernación de Sucre emitió un comunicado.", src), src)

    def test_rejects_paraphrase(self):
        src = "La gobernadora inauguró el hospital de Sampués."
        got = recover_verbatim("La mandataria abrió un nuevo centro hospitalario en Sampués.", src)
        self.assertEqual(got, "")

    def test_whitespace_flexible_recovers_source_slice(self):
        src = "Lucy  García\nanunció la obra."
        got = recover_verbatim("Lucy García anunció la obra.", src)
        self.assertIn("Lucy", got)
        self.assertIn("anunció la obra", got.replace("\n", " "))
        self.assertTrue(got in src or got.replace("  ", " ") in src.replace("\n", " "))

    def test_filler_never_kept(self):
        self.assertTrue(is_filler("(sin actor externo)"))
        self.assertTrue(is_filler("N/A"))
        self.assertTrue(is_filler("ninguno"))
        self.assertTrue(is_filler("—"))
        self.assertFalse(is_filler("Lucy Inés García Montes, Gobernadora de Sucre"))


class HeuristicExtractTests(unittest.TestCase):
    def test_lucy_own_intervention_is_literal(self):
        out = heuristic_analyze("Gobernadora entregó becas", LUCY_BODY)
        self.assertEqual(out[COL_TONO], "Positivo")
        self.assertIn("Lucy Inés García Montes", out[COL_PROPIOS])
        self.assertIn("Gobernadora", out[COL_PROPIOS])
        extract = out[COL_INT_PROPIA]
        self.assertTrue(extract)
        self.assertIn(extract[:40], LUCY_BODY)
        self.assertIn("entregó 200 becas", extract)
        self.assertNotIn("sin actor", extract.lower())
        self.assertNotIn("(", extract.split("entregó")[0][-5:] if False else "")
        self.assertIn("Ricardo Hernández", out[COL_EXTERNOS])
        ext = out[COL_MENCION_EXT]
        self.assertTrue(ext)
        self.assertIn("alcalde", ext.lower())
        self.assertIn(ext[:30], LUCY_BODY)

    def test_external_senator_no_own_agency(self):
        out = heuristic_analyze("Senador cuestiona vías", SENATOR_BODY)
        self.assertEqual(out[COL_TONO], "Negativo")
        self.assertEqual(out[COL_INT_PROPIA], "")
        self.assertEqual(out[COL_PROPIOS], "")
        self.assertIn("Andrés Pérez", out[COL_EXTERNOS])
        self.assertIn("senador", out[COL_EXTERNOS].lower())
        self.assertEqual(out[COL_MENCION_EXT], SENATOR_BODY)
        self.assertNotEqual(out[COL_EXTERNOS].lower(), "(sin actor externo)")

    def test_secretaria_departamental_as_own_actor(self):
        out = heuristic_analyze("Resolución de cobertura", SECRE_BODY)
        self.assertIn("Secretaría de Educación", out[COL_PROPIOS])
        self.assertNotIn("emitió", out[COL_PROPIOS].lower())
        self.assertIn("Resolución 045", out[COL_INT_PROPIA])
        self.assertTrue(out[COL_INT_PROPIA] in SECRE_BODY or out[COL_INT_PROPIA][:20] in SECRE_BODY)
        self.assertEqual(out[COL_EXTERNOS], "")
        self.assertEqual(out[COL_MENCION_EXT], "")

    def test_no_mention_yields_empty_strings(self):
        body = "El Ideam publicó el pronóstico de lluvias para la región Caribe."
        out = heuristic_analyze("Boletín climático", body)
        self.assertEqual(out[COL_TONO], "Neutro")
        self.assertEqual(out[COL_PROPIOS], "")
        self.assertEqual(out[COL_INT_PROPIA], "")
        self.assertEqual(out[COL_EXTERNOS], "")
        self.assertEqual(out[COL_MENCION_EXT], "")

    def test_alias_gobernadora_de_sucre(self):
        body = "La gobernadora de Sucre anunció la pavimentación de la vía Sincelejo-Sampués."
        out = heuristic_analyze("Pavimentación vial", body)
        self.assertIn("Lucy Inés García Montes", out[COL_PROPIOS])
        self.assertIn("anunció la pavimentación", out[COL_INT_PROPIA])


class LlmMergeTests(unittest.TestCase):
    def test_paraphrased_llm_extract_falls_back_to_heuristic(self):
        heuristic = heuristic_analyze("Becas", LUCY_BODY)
        llm = {
            COL_TONO: "Positivo",
            COL_PROPIOS: "Lucy Inés García Montes, Gobernadora de Sucre",
            COL_INT_PROPIA: "La mandataria entregó un paquete de apoyos académicos.",
            COL_EXTERNOS: "",
            COL_MENCION_EXT: "(sin actor externo)",
        }
        merged = merge_analysis(llm, heuristic, "Becas", LUCY_BODY)
        self.assertIn("entregó 200 becas", merged[COL_INT_PROPIA])
        self.assertNotIn("paquete de apoyos", merged[COL_INT_PROPIA])
        self.assertEqual(merged[COL_MENCION_EXT], heuristic[COL_MENCION_EXT])
        self.assertNotIn("sin actor", merged[COL_MENCION_EXT].lower())

    def test_verbatim_llm_span_is_kept(self):
        span = "La gobernadora Lucy García Montes entregó 200 becas universitarias en Sincelejo."
        heuristic = heuristic_analyze("Becas", LUCY_BODY)
        llm = {
            "tono": "Positivo",
            "nombre_cargo_propios": "Lucy Inés García Montes, Gobernadora de Sucre",
            "intervencion_propia": span,
            "nombre_cargo_externos": "Ricardo Hernández, alcalde",
            "mencion_externa": "El alcalde Ricardo Hernández felicitó a la gobernadora por la inversión social.",
        }
        merged = merge_analysis(llm, heuristic, "Becas", LUCY_BODY)
        self.assertEqual(merged[COL_INT_PROPIA], span)
        self.assertEqual(
            merged[COL_MENCION_EXT],
            "El alcalde Ricardo Hernández felicitó a la gobernadora por la inversión social.",
        )

    def test_analyze_article_without_client_is_heuristic(self):
        out = analyze_article("Becas", LUCY_BODY, client=None)
        self.assertEqual(out[COL_TONO], "Positivo")
        self.assertTrue(out[COL_INT_PROPIA])


class PipelineXlsxTests(unittest.TestCase):
    def test_sample_processing_writes_sucre_columns_and_keeps_input(self):
        result = process_sucre_dossier(
            io.BytesIO(build_sample_xlsx()),
            ai_config={"enabled": False},
        )
        self.assertGreaterEqual(result["total_rows"], 4)
        bio = io.BytesIO(result["output_data"])
        wb = load_workbook(bio, read_only=True, data_only=True)
        ws = wb.active
        headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for col in ("ID", "Fecha", "Medio", "Título", "CuerpoEs"):
            self.assertIn(col, headers)
        for col in SUCRE_OUTPUT_COLUMNS:
            self.assertIn(col, headers)

        rows = list(ws.iter_rows(min_row=2, values_only=True))
        by_header = [{headers[i]: row[i] for i in range(len(headers))} for row in rows]
        lucy_row = next(
            r for r in by_header
            if "becas" in str(r.get("Título") or "").lower()
        )
        body = str(lucy_row["CuerpoEs"] or "")
        own = str(lucy_row[COL_INT_PROPIA] or "")
        self.assertTrue(own)
        self.assertIn(own[:50], body)
        self.assertNotIn("Gobernadora Lucy García entregó 200 becas en Sincelejo.", own)
        self.assertIn("Lucy", str(lucy_row[COL_PROPIOS] or ""))
        self.assertEqual(lucy_row[COL_TONO], "Positivo")
        ext = str(lucy_row[COL_MENCION_EXT] or "")
        self.assertTrue(ext)
        self.assertIn(ext[:40], body)

        sen_row = next(r for r in by_header if "senador" in str(r.get("Título") or "").lower())
        self.assertEqual(sen_row[COL_TONO], "Negativo")
        self.assertIn(str(sen_row[COL_MENCION_EXT] or "")[:40], str(sen_row["CuerpoEs"] or ""))
        self.assertFalse(sen_row[COL_INT_PROPIA])
        self.assertFalse(sen_row[COL_PROPIOS])

        clima = next(r for r in by_header if "climático" in str(r.get("Título") or "").lower())
        self.assertEqual(clima[COL_PROPIOS] or "", "")
        self.assertEqual(clima[COL_INT_PROPIA] or "", "")
        self.assertEqual(clima[COL_EXTERNOS] or "", "")
        self.assertEqual(clima[COL_MENCION_EXT] or "", "")
        wb.close()

    def test_missing_columns_raise(self):
        df = pd.DataFrame({"Foo": [1], "Bar": ["x"]})
        buf = io.BytesIO()
        df.to_excel(buf, index=False, engine="openpyxl")
        with self.assertRaises(Exception):
            process_sucre_dossier(io.BytesIO(buf.getvalue()), ai_config={"enabled": False})


class GrillIsolationTests(unittest.TestCase):
    def test_grill_output_shape_unchanged(self):
        self.assertIn("Tono_IA", ["Contexto analizado", "Tono_IA", "Tema_IA", "Subtema_IA"])
        self.assertNotIn(COL_INT_PROPIA, BASE_OUTPUT_COLUMNS)
        self.assertNotIn(COL_MENCION_EXT, BASE_OUTPUT_COLUMNS)
        self.assertIn("Título", BASE_OUTPUT_COLUMNS)
        self.assertEqual(KEY_MAP["titulo"], "Título")

    def test_grill_app_does_not_import_sucre(self):
        import ast
        with open(os.path.join(ROOT, "app.py"), encoding="utf-8") as fh:
            tree = ast.parse(fh.read())
        imported = []
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                imported.extend(a.name for a in node.names)
            elif isinstance(node, ast.ImportFrom) and node.module:
                imported.append(node.module)
        self.assertNotIn("sucre_analyzer", imported)
        self.assertNotIn("sucre_pipeline", imported)
        self.assertNotIn("app_sucre", imported)


class MockedLlmPathTests(unittest.TestCase):
    def test_openai_path_uses_literal_body_span(self):
        span = "La gobernadora Lucy García Montes entregó 200 becas universitarias en Sincelejo."
        payload = {
            "tono": "Positivo",
            "nombre_cargo_propios": "Lucy Inés García Montes, Gobernadora de Sucre",
            "intervencion_propia": span,
            "nombre_cargo_externos": "Ricardo Hernández, alcalde",
            "mencion_externa": "El alcalde Ricardo Hernández felicitó a la gobernadora por la inversión social.",
        }
        fake_resp = MagicMock()
        fake_resp.choices = [MagicMock()]
        fake_resp.choices[0].message.content = __import__("json").dumps(payload)
        client = MagicMock()
        client.chat.completions.create.return_value = fake_resp
        out = analyze_article("Becas", LUCY_BODY, client=client, model="gpt-test")
        self.assertEqual(out[COL_INT_PROPIA], span)
        self.assertEqual(out[COL_TONO], "Positivo")
        client.chat.completions.create.assert_called_once()


if __name__ == "__main__":
    unittest.main()

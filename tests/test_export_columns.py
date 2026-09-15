# ======================================
# Contrato de columnas Excel: sin resumen corto / revalorización,
# Resumen - Aclaracion termina en un punto,
# Tono/Tema/Subtema después de Audiencia, Contexto analizado al final.
# ======================================
import io
import os
import sys
import unittest

from openpyxl import load_workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from pipeline import (  # noqa: E402
    AI_COLUMNS_AFTER_AUDIENCIA,
    BASE_OUTPUT_COLUMNS,
    CONTEXTO_COLUMN,
    KEY_MAP,
    build_export_columns,
    corregir_texto,
    generate_output_excel,
)
from sucre_analyzer import SUCRE_ACTOR_COLUMNS  # noqa: E402

EXPECTED_BASE = [
    "ID Noticia", "Fecha", "Hora", "Medio", "Tipo de Medio",
    "Sección - Programa", "Región", "Título", "Autor - Conductor",
    "Nro. Pagina", "Dimensión", "Duración - Nro. Caracteres",
    "CPE", "Tier", "Audiencia",
    "Link Nota", "Resumen - Aclaracion", "Link (Streaming - Imagen)", "Menciones - Empresa",
    "ID duplicada",
]

EXPECTED_WITH_AI = [
    "ID Noticia", "Fecha", "Hora", "Medio", "Tipo de Medio",
    "Sección - Programa", "Región", "Título", "Autor - Conductor",
    "Nro. Pagina", "Dimensión", "Duración - Nro. Caracteres",
    "CPE", "Tier", "Audiencia",
    "Tono_IA", "Tema_IA", "Subtema_IA",
    "Link Nota", "Resumen - Aclaracion", "Link (Streaming - Imagen)", "Menciones - Empresa",
    "ID duplicada",
    "Contexto analizado",
]

EXPECTED_SUCRE_WITH_AI = (
    EXPECTED_WITH_AI[:-1] + list(SUCRE_ACTOR_COLUMNS) + [EXPECTED_WITH_AI[-1]]
)
EXPECTED_SUCRE_WITHOUT_AI = EXPECTED_BASE + list(SUCRE_ACTOR_COLUMNS)


class BaseOutputColumnsTests(unittest.TestCase):
    def test_dropped_columns_are_not_exported(self):
        self.assertNotIn("resumen corto", BASE_OUTPUT_COLUMNS)
        self.assertNotIn("revalorización", BASE_OUTPUT_COLUMNS)
        self.assertEqual(list(BASE_OUTPUT_COLUMNS), EXPECTED_BASE)

    def test_ai_insert_columns_are_tono_tema_subtema_only(self):
        self.assertEqual(
            list(AI_COLUMNS_AFTER_AUDIENCIA),
            ["Tono_IA", "Tema_IA", "Subtema_IA"],
        )
        self.assertNotIn(CONTEXTO_COLUMN, AI_COLUMNS_AFTER_AUDIENCIA)
        self.assertEqual(CONTEXTO_COLUMN, "Contexto analizado")

    def test_generate_output_excel_omits_dropped_columns(self):
        rows = [
            {
                "ID Noticia": 1,
                "Título": "Nota",
                "resumen corto": "no debe salir",
                "revalorización": 123,
                "Resumen - Aclaracion": "Cuerpo de la nota.",
                "Contexto analizado": "marca en el párrafo",
                "Tono_IA": "Neutro",
                "Tema_IA": "Tema",
                "Subtema_IA": "Subtema",
            }
        ]
        data = generate_output_excel(rows, KEY_MAP)
        wb = load_workbook(io.BytesIO(data))
        headers = [c.value for c in wb["Resultado"][1]]
        self.assertNotIn("resumen corto", headers)
        self.assertNotIn("revalorización", headers)
        self.assertEqual(headers, EXPECTED_BASE)
        wb.close()

        data_ai = generate_output_excel(
            rows, KEY_MAP, columns_to_use=build_export_columns(True)
        )
        wb = load_workbook(io.BytesIO(data_ai))
        headers = [c.value for c in wb["Resultado"][1]]
        self.assertEqual(headers, EXPECTED_WITH_AI)
        self.assertEqual(headers[-1], "Contexto analizado")
        self.assertNotIn("resumen corto", headers)
        self.assertNotIn("revalorización", headers)
        wb.close()

    def test_build_export_columns_full_order(self):
        self.assertEqual(build_export_columns(False), EXPECTED_BASE)
        with_ai = build_export_columns(True)
        self.assertEqual(with_ai, EXPECTED_WITH_AI)
        audiencia_i = with_ai.index("Audiencia")
        self.assertEqual(with_ai[audiencia_i + 1:audiencia_i + 4], ["Tono_IA", "Tema_IA", "Subtema_IA"])
        self.assertEqual(with_ai[audiencia_i + 4], "Link Nota")
        self.assertEqual(with_ai[-1], "Contexto analizado")

        self.assertEqual(build_export_columns(False, SUCRE_ACTOR_COLUMNS), EXPECTED_SUCRE_WITHOUT_AI)
        sucre = build_export_columns(True, SUCRE_ACTOR_COLUMNS)
        self.assertEqual(sucre, EXPECTED_SUCRE_WITH_AI)
        audiencia_i = sucre.index("Audiencia")
        self.assertEqual(sucre[audiencia_i + 1:audiencia_i + 4], ["Tono_IA", "Tema_IA", "Subtema_IA"])
        self.assertEqual(sucre[audiencia_i + 4], "Link Nota")
        self.assertEqual(sucre[-1], "Contexto analizado")
        for col in SUCRE_ACTOR_COLUMNS:
            self.assertLess(sucre.index(col), sucre.index("Contexto analizado"))
            self.assertGreater(sucre.index(col), sucre.index("ID duplicada"))
        self.assertNotIn("resumen corto", sucre)
        self.assertNotIn("revalorización", sucre)


class CorregirTextoTests(unittest.TestCase):
    def test_ends_with_single_period_not_ellipsis(self):
        self.assertEqual(corregir_texto("La marca anunció un plan"), "La marca anunció un plan.")
        self.assertEqual(corregir_texto("La marca anunció un plan."), "La marca anunció un plan.")
        self.assertEqual(corregir_texto("La marca anunció un plan..."), "La marca anunció un plan.")

    def test_empty_stays_empty(self):
        self.assertEqual(corregir_texto(""), "")
        self.assertEqual(corregir_texto("nan"), "")


if __name__ == "__main__":
    unittest.main()
